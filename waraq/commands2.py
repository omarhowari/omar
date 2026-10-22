"""Commands added for the PDF-XChange-style tabs (View, Comment, Protect, Form, Document, Convert, Review,
Accessibility, Bookmarks, Help). Mixed into MainWindow."""
import datetime as _dt
import html as _html
import json
import os
import re
import subprocess
import sys
import tempfile

import pymupdf as fitz
from PySide6.QtCore import Qt, QUrl
from PySide6.QtGui import QDesktopServices, QImage
from PySide6.QtWidgets import (QApplication, QFileDialog, QInputDialog, QMenu, QMessageBox)

from . import ops, convert, i18n
from .dialogs import ask, ReportDialog, PAGES_HELP
from .document import PdfDocument
from .i18n import tr
from .icons import icon, cicon


def A(en, ar):
    return ar if i18n.LANG == "ar" else en


class Commands2Mixin:
    # ------------------------------------------------------------------ menus
    def _other_tools_menu(self):
        m = QMenu(self)
        for c, l, i in (("snapshot", A("Snapshot", "لقطة شاشة"), "snapshot"), ("zoom_tool", A("Zoom tool", "أداة التكبير"), "zoomin"),
                        ("loupe", A("Loupe", "مكبرة"), "search"), ("eraser", A("Eraser", "ممحاة"), "eraser"),
                        ("select_comments", A("Edit comments", "تحرير التعليقات"), "selectcomment"),
                        ("edittext", A("Edit text", "تحرير النص"), "typewriter"),
                        ("measure", A("Distance", "مسافة"), "line")):
            m.addAction(self._action(c, l, i))
        return m

    def _rotateview_menu(self):
        m = QMenu(self)
        m.addAction(self._action("rotate_view_right", A("Rotate view right", "تدوير العرض لليمين"), "rotate"))
        m.addAction(self._action("rotate_view_left", A("Rotate view left", "تدوير العرض لليسار"), "rotateleft"))
        m.addAction(self._action("rotate_view_reset", A("Reset view rotation", "إعادة ضبط التدوير"), "undo"))
        return m

    def _panes_menu(self):
        m = QMenu(self)
        for c, l, i in (("panel_thumbs", A("Pages", "الصفحات"), "thumbs"), ("panel_bookmarks", A("Bookmarks", "الإشارات المرجعية"), "bookmark"),
                        ("panel_comments", A("Comments", "التعليقات"), "comments"), ("panel_search", A("Search", "بحث"), "search"),
                        ("panel_attach", A("Attachments", "المرفقات"), "attach"), ("toggle_props", A("Properties", "خصائص"), "properties")):
            m.addAction(self._action(c, l, i))
        return m

    def _show_menu(self):
        m = QMenu(self)
        m.addAction(A("Show all comments", "إظهار كل التعليقات"), lambda: self._set_comments_visible(True))
        m.addAction(A("Hide all comments", "إخفاء كل التعليقات"), lambda: self._set_comments_visible(False))
        return m

    def _set_comments_visible(self, on):
        self.actions["hide_comments"].setChecked(not on)
        self.do_hide_comments()

    # ------------------------------------------------------------------ quick launch / search
    def quick_launch(self, text):
        text = (text or "").strip().lower()
        if not text:
            return
        for cmd, a in self.actions.items():
            if a.text().lower() == text:
                a.trigger()
                return
        for cmd, a in self.actions.items():
            if text in a.text().lower():
                a.trigger()
                return
        self.statusBar().showMessage(A("Command not found", "لم يتم العثور على الأمر"), 3000)

    def do_quick_find(self):
        t = self._need()
        if not t:
            return
        q, ok = QInputDialog.getText(self, A("Quick search", "بحث سريع"), A("Find:", "بحث عن:"), text=t.view.selected_text())
        if ok and q:
            t.show_panel(3)
            t.search.edit.setText(q)
            t.search.run()

    def _goto_from_box(self, box):
        t = self.tab()
        if not t:
            return
        m = re.match(r"\s*(\d+)", box.text())
        if m:
            t.view.go_to(int(m.group(1)) - 1)

    # ------------------------------------------------------------------ view
    def do_fitvisible(self):
        t = self._need()
        if not t:
            return
        page = t.doc[t.view.current]
        bbox = fitz.Rect()
        for b in page.get_text("blocks"):
            bbox |= fitz.Rect(b[:4])
        for d in page.get_drawings():
            bbox |= d["rect"]
        for im in page.get_image_info():
            bbox |= fitz.Rect(im["bbox"])
        if bbox.is_empty:
            return t.view.set_fit("width")
        bbox = (bbox + (-12, -12, 12, 12)) & page.rect
        from .viewer import BASE
        vw = t.view.viewport().width() - 40
        t.view.set_zoom(vw / (bbox.width * BASE))
        t.view.go_to_rect(t.view.current, bbox * page.derotation_matrix)

    def _rot_view(self, d):
        t = self._need()
        if t:
            t.view.set_view_rotation(0 if d is None else t.view.view_rotation + d)

    def do_rotate_view_right(self):
        self._rot_view(90)

    def do_rotate_view_left(self):
        self._rot_view(-90)

    def do_rotate_view_reset(self):
        self._rot_view(None)

    def do_layout_single2(self):
        self._layout("single2")

    def _toggle_view_flag(self, cmd, attr):
        on = self.actions[cmd].isChecked()
        for t in self.doc_tabs():
            setattr(t.view, attr, on)
            t.view.viewport().update()

    def do_rulers(self):
        self._toggle_view_flag("rulers", "show_rulers")

    def do_grid(self):
        self._toggle_view_flag("grid", "show_grid")

    def do_hand_text(self):
        self._toggle_view_flag("hand_text", "hand_text")

    def do_exclusive(self):
        self.exclusive = self.actions["exclusive"].isChecked()
        self.settings.setValue("exclusive", self.exclusive)

    def do_smooth_scroll(self):
        on = self.actions["smooth_scroll"].isChecked()
        for t in self.doc_tabs():
            t.view.verticalScrollBar().setSingleStep(8 if on else 40)

    def do_doc_tabs(self):
        self.tabs.tabBar().setVisible(self.actions["doc_tabs"].isChecked())

    def do_recent(self):
        self.do_open()

    def do_panes(self):
        self._panel(0)

    def do_other_tools(self):
        self.set_tool("snapshot")

    def do_show_comments(self):
        self._set_comments_visible(self.actions["hide_comments"].isChecked())

    def do_highlight_fields(self):
        on = self.actions["highlight_fields"].isChecked()
        t = self._need()
        if not t:
            return
        t.pdf.checkpoint()
        for p in t.doc:
            for w in p.widgets():
                w.fill_color = (0.86, 0.9, 1.0) if on else None
                w.update()
        t.done(-1)

    # ------------------------------------------------------------------ comments navigation
    def _comment_list(self, t):
        out = []
        for i in range(t.doc.page_count):
            pg = t.doc[i]
            for a in pg.annots():
                if a.type[1] not in ("Popup", "Link", "Widget"):
                    out.append((i, a.xref, a.rect.y0))
        return out

    def _step_comment(self, d):
        t = self._need()
        if not t:
            return
        lst = self._comment_list(t)
        if not lst:
            self.statusBar().showMessage(A("No comments", "لا توجد تعليقات"), 3000)
            return
        cur = t.view.sel_annot
        k = next((n for n, (p, x, _) in enumerate(lst) if cur and (p, x) == cur), -1 if d > 0 else 0)
        p, x, _ = lst[(k + d) % len(lst)]
        t.goto_annot(p, x)

    def do_next_comment(self):
        self._step_comment(1)

    def do_prev_comment(self):
        self._step_comment(-1)

    # ------------------------------------------------------------------ protect
    def refresh_signatures(self):
        t = self.tab()
        names = []
        if t:
            for p in t.doc:
                for w in p.widgets():
                    if w.field_type == fitz.PDF_WIDGET_TYPE_SIGNATURE:
                        names.append(w.field_name + (" ✔" if w.is_signed else ""))
        for lw in getattr(self, "sig_lists", []):
            lw.clear()
            lw.addItems(names or [A("No signatures", "لا توجد توقيعات")])

    def do_invisible_sign(self):
        t = self._need()
        if t:
            self.digital_sign(t.view.current, None)

    def do_timestamp(self):
        t = self._need()
        if not t:
            return
        if t.pdf.is_modified() or not t.pdf.path:
            if not self.save(t):
                return
        v = ask(self, A("Timestamp", "الطابع الزمني"), [
            ("tsa", A("Timestamp server", "خادم الطابع الزمني"), "text", self.settings.value("tsa", "") or "http://timestamp.digicert.com"),
            ("out", A("Save to", "حفظ إلى"), "file_save", os.path.splitext(t.pdf.path)[0] + "_timestamped.pdf", "PDF (*.pdf)")],
            help_text=A("Adds an RFC-3161 document timestamp. Needs internet.", "يضيف ختماً زمنياً معتمداً للمستند. يحتاج اتصال إنترنت."))
        if not v or not v["out"]:
            return
        from pyhanko.sign import timestamps, signers
        from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            with open(t.pdf.path, "rb") as inf:
                w = IncrementalPdfFileWriter(inf, strict=False)
                ts = timestamps.HTTPTimeStamper(v["tsa"])
                with open(v["out"], "wb") as outf:
                    signers.PdfTimeStamper(ts).timestamp_pdf(w, "sha256", output=outf)
        finally:
            QApplication.restoreOverrideCursor()
        self.settings.setValue("tsa", v["tsa"])
        self.open_files([v["out"]])

    def do_remove_signatures(self):
        t = self._need()
        if not t:
            return
        n = sum(1 for p in t.doc for w in p.widgets() if w.field_type == fitz.PDF_WIDGET_TYPE_SIGNATURE)
        hand = sum(1 for p in t.doc for a in p.annots() if a.info.get("subject") == "Signature")
        if not n and not hand:
            QMessageBox.information(self, A("Signatures", "التوقيعات"), A("No signatures", "لا توجد توقيعات"))
            return
        if QMessageBox.question(self, A("Signatures", "التوقيعات"),
                                A(f"Remove {n + hand} signature(s)?", f"حذف {n + hand} توقيع؟")) != QMessageBox.Yes:
            return
        t.pdf.checkpoint()
        for p in t.doc:
            for w in list(p.widgets()):
                if w.field_type == fitz.PDF_WIDGET_TYPE_SIGNATURE:
                    p.delete_widget(w)
            for a in list(p.annots()):
                if a.info.get("subject") == "Signature":
                    p.delete_annot(a)
        t.done(-1)
        self.refresh_signatures()

    def do_delete_hand_signatures(self):
        t = self._need()
        if not t:
            return
        t.pdf.checkpoint()
        n = 0
        for p in t.doc:
            for a in list(p.annots()):
                if a.info.get("subject") == "Signature":
                    p.delete_annot(a)
                    n += 1
        t.done(-1, f"{n}")

    def do_sign_all_pages(self):
        """Copy the selected (or last) hand signature to the same spot on every page."""
        t = self._need()
        if not t:
            return
        src = None
        if t.view.sel_annot:
            src = ops.load_annot(t.doc, *t.view.sel_annot)
        if src is None or src.info.get("subject") != "Signature":
            for i in range(t.doc.page_count):
                pg = t.doc[i]
                for a in pg.annots():
                    if a.info.get("subject") == "Signature":
                        src = a
                        src._page_ref = pg
        if src is None:
            QMessageBox.information(self, A("Signature", "توقيع"), A("Place a signature first (Protect › Manage).",
                                                                     "ضع توقيعاً أولاً (حماية › إدارة)."))
            return
        rect = fitz.Rect(src.rect)
        sp = src._page_ref.number
        pix = src._page_ref.get_pixmap(clip=rect * src._page_ref.rotation_matrix, dpi=220, alpha=True, annots=True)
        t.pdf.checkpoint()
        for i in range(t.doc.page_count):
            if i == sp:
                continue
            pg = t.doc[i]
            if (rect & pg.mediabox).is_empty:
                continue
            a = pg.add_stamp_annot(rect, stamp=t.doc[sp].get_pixmap(clip=rect, dpi=220, alpha=True) if False else pix)
            a.set_info(title=self.style["author"], subject="Signature", content="Signature")
            a.update()
        t.done(-1, A("Signature applied to all pages", "تم تطبيق التوقيع على كل الصفحات"))

    def do_clear_redactions(self):
        t = self._need()
        if not t:
            return
        t.pdf.checkpoint()
        for p in t.doc:
            for a in list(p.annots(types=[fitz.PDF_ANNOT_REDACT])):
                p.delete_annot(a)
        t.done(-1)

    # ------------------------------------------------------------------ forms
    def do_fill_forms(self):
        self.set_tool("select")
        self.statusBar().showMessage(A("Click a field to fill it", "انقر على حقل لتعبئته"), 4000)

    def do_email_form(self):
        t = self._need()
        if not t:
            return
        data = t.form_data()
        fn = os.path.join(tempfile.mkdtemp(), os.path.splitext(t.pdf.title)[0] + "_data.json")
        json.dump(data, open(fn, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        self._email_file(fn)

    # ------------------------------------------------------------------ document structure
    def do_paste_page(self):
        img = QApplication.clipboard().image()
        t = self._need()
        if not t or img.isNull():
            self.statusBar().showMessage(tr("No image in clipboard"), 3000)
            return
        tmp = os.path.join(tempfile.mkdtemp(), "clip.png")
        img.save(tmp)
        t.pdf.checkpoint()
        t.doc.insert_pdf(ops.images_to_pdf([tmp], "A4", 10), start_at=t.view.current + 1)
        t.done(-1)

    def do_rotate_180(self):
        t = self._need()
        if t:
            pages, _ = self._ask_pages(A("Rotate 180°", "تدوير 180°"))
            if pages:
                self.rotate_pages(180, pages)

    def do_replace_pages(self):
        t = self._need()
        if not t:
            return
        fn, _ = QFileDialog.getOpenFileName(self, A("Replace pages", "استبدل الصفحات"), "", "PDF (*.pdf)")
        if not fn:
            return
        src = ops.open_any(fn)
        v = ask(self, A("Replace pages", "استبدل الصفحات"), [
            ("pages", A("Replace pages (this document)", "الصفحات المراد استبدالها"), "text", str(t.view.current + 1)),
            ("from", A("With pages (other file)", "بالصفحات من الملف الآخر"), "text", "1")], help_text=PAGES_HELP)
        if not v:
            return
        dst = ops.parse_ranges(v["pages"], t.doc.page_count)
        srcp = ops.parse_ranges(v["from"], src.page_count)
        if not dst or not srcp:
            return
        t.pdf.checkpoint()
        at = dst[0]
        t.doc.delete_pages(dst)
        for k, sp in enumerate(srcp):
            t.doc.insert_pdf(src, from_page=sp, to_page=sp, start_at=at + k)
        t.done(-1, A("Pages replaced", "تم استبدال الصفحات"))

    def do_swap_pages(self):
        t = self._need()
        if not t:
            return
        v = ask(self, A("Swap pages", "تبادل الصفحات"), [("a", A("Page", "الصفحة"), "int", t.view.current + 1, (1, t.doc.page_count)),
                                                        ("b", A("With page", "مع الصفحة"), "int", min(t.view.current + 2, t.doc.page_count),
                                                         (1, t.doc.page_count))])
        if not v or v["a"] == v["b"]:
            return
        order = list(range(t.doc.page_count))
        a, b = v["a"] - 1, v["b"] - 1
        order[a], order[b] = order[b], order[a]
        t.pdf.checkpoint()
        t.doc.select(order)
        t.done(-1)

    def do_overlay(self):
        t = self._need()
        if not t:
            return
        fn, _ = QFileDialog.getOpenFileName(self, A("Overlay", "تراكب"), "", "PDF (*.pdf)")
        if not fn:
            return
        src = ops.open_any(fn)
        v = ask(self, A("Overlay", "تراكب"), [("page", A("Page of the other file", "صفحة الملف الآخر"), "int", 1, (1, src.page_count)),
                                              ("pages", A("Apply to pages", "تطبيق على الصفحات"), "text", "all"),
                                              ("behind", A("Behind content (underlay)", "خلف المحتوى"), "bool", False)])
        if not v:
            return
        t.pdf.checkpoint()
        for i in ops.parse_ranges(v["pages"], t.doc.page_count):
            pg = t.doc[i]
            pg.show_pdf_page(pg.mediabox, src, v["page"] - 1, overlay=not v["behind"])
        t.done(-1)

    def do_page_labels(self):
        t = self._need()
        if not t:
            return
        v = ask(self, A("Number pages", "ترقيم الصفحات (التسميات)"), [
            ("start", A("From page", "من الصفحة"), "int", 1, (1, t.doc.page_count)),
            ("style", A("Style", "النمط"), "choice", "D", [("D", "1, 2, 3"), ("r", "i, ii, iii"), ("R", "I, II, III"),
                                                          ("a", "a, b, c"), ("A", "A, B, C")]),
            ("prefix", A("Prefix", "بادئة"), "text", ""), ("first", A("Start number", "رقم البداية"), "int", 1, (1, 100000))])
        if not v:
            return
        labels = [l for l in (t.doc.get_page_labels() or []) if l["startpage"] != v["start"] - 1]
        labels.append({"startpage": v["start"] - 1, "prefix": v["prefix"], "style": v["style"], "firstpagenum": v["first"]})
        t.pdf.checkpoint()
        t.doc.set_page_labels(sorted(labels, key=lambda l: l["startpage"]))
        t.done(-2, A("Page labels set", "تم ضبط تسميات الصفحات"))

    # ------------------------------------------------------------------ convert
    def do_new_blank(self):
        v = ask(self, A("New document", "مستند جديد"), [("size", A("Paper", "الورق"), "choice", "A4", list(ops.PAPER)),
                                                         ("n", A("Pages", "الصفحات"), "int", 1, (1, 500)),
                                                         ("land", A("Landscape", "أفقي"), "bool", False)])
        if not v:
            return
        w, h = ops.PAPER[v["size"]]
        if v["land"]:
            w, h = h, w
        d = fitz.open()
        for _ in range(v["n"]):
            d.new_page(width=w, height=h)
        pdf = PdfDocument(doc=d)
        pdf.suggested = A("New document.pdf", "مستند جديد.pdf")
        pdf.set_modified(True)
        self.add_tab(pdf)

    def do_rasterize(self):
        t = self._need()
        if not t:
            return
        v = ask(self, A("Content to images", "تحويل المحتوى إلى صور"), [("dpi", "dpi", "choice", 200, [100, 150, 200, 300]),
                                                                        ("pages", tr("Pages to process"), "text", "all")])
        if not v:
            return
        pages = set(ops.parse_ranges(v["pages"], t.doc.page_count))
        out = fitz.open()
        for i in range(t.doc.page_count):
            pg = t.doc[i]
            if i in pages:
                pm = pg.get_pixmap(dpi=v["dpi"], alpha=False)
                np_ = out.new_page(width=pg.rect.width, height=pg.rect.height)
                np_.insert_image(np_.rect, pixmap=pm)
            else:
                out.insert_pdf(t.doc, from_page=i, to_page=i)
        self._replace_doc(t, out, A("Pages converted to images", "تم تحويل الصفحات إلى صور"))

    def _render_doc(self, t, fn_pix, dpi=200):
        out = fitz.open()
        for pg in t.doc:
            pm = fn_pix(pg.get_pixmap(dpi=dpi, alpha=False))
            np_ = out.new_page(width=pg.rect.width, height=pg.rect.height)
            np_.insert_image(np_.rect, pixmap=pm)
        try:
            out.set_toc(t.doc.get_toc())
        except Exception:
            pass
        return out

    def do_grayscale(self):
        t = self._need()
        if not t:
            return
        t.pdf.checkpoint()
        try:
            t.doc.recolor(1)
            t.done(-1, A("Converted to grayscale", "تم التحويل إلى تدرج رمادي"))
        except Exception:
            out = self._render_doc(t, lambda pm: fitz.Pixmap(fitz.csGRAY, pm))
            old = t.pdf.doc
            t.pdf.doc = out
            old.close()
            t.done(-1)

    def do_enhance(self):
        t = self._need()
        if not t:
            return
        v = ask(self, A("Enhance scanned pages", "تحسين الصفحات الممسوحة ضوئياً"), [
            ("gray", A("Convert to grayscale", "تحويل إلى رمادي"), "bool", True),
            ("gamma", A("Darken text (gamma)", "تغميق النص"), "float", 0.8, (0.3, 2.0)),
            ("dpi", "dpi", "choice", 200, [150, 200, 300])])
        if not v:
            return

        def fx(pm):
            if v["gray"]:
                pm = fitz.Pixmap(fitz.csGRAY, pm)
            pm.gamma_with(v["gamma"])
            return pm
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            out = self._render_doc(t, fx, v["dpi"])
        finally:
            QApplication.restoreOverrideCursor()
        self._replace_doc(t, out, A("Pages enhanced", "تم تحسين الصفحات"))

    def do_to_ppt(self):
        t = self._need()
        if not t:
            return
        fn, _ = QFileDialog.getSaveFileName(self, "PowerPoint", os.path.splitext(t.pdf.path or t.pdf.title)[0] + ".pptx",
                                            "PowerPoint (*.pptx)")
        if not fn:
            return
        try:
            from pptx import Presentation
            from pptx.util import Emu
        except ImportError:
            QMessageBox.warning(self, "PowerPoint", "pip install python-pptx")
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            prs = Presentation()
            p0 = t.doc[0].rect
            prs.slide_width = Emu(int(p0.width * 12700))
            prs.slide_height = Emu(int(p0.height * 12700))
            tmp = tempfile.mkdtemp()
            for i, pg in enumerate(t.doc):
                img = os.path.join(tmp, f"p{i}.png")
                pg.get_pixmap(dpi=150, alpha=False).save(img)
                s = prs.slides.add_slide(prs.slide_layouts[6])
                s.shapes.add_picture(img, 0, 0, prs.slide_width, prs.slide_height)
                notes = pg.get_text().strip()
                if notes:
                    s.notes_slide.notes_text_frame.text = notes[:5000]
            prs.save(fn)
        finally:
            QApplication.restoreOverrideCursor()
        QDesktopServices.openUrl(QUrl.fromLocalFile(fn))

    def do_to_excel(self):
        t = self._need()
        if not t:
            return
        from . import pdfconv
        v = ask(self, A("Convert to Excel", "تحويل إلى Excel"), [
            ("layout", A("Sheets", "الأوراق"), "choice", "per_table",
             [("per_table", A("One sheet per table (continued tables joined)", "ورقة لكل جدول (مع دمج الجداول المتصلة)")),
              ("per_page", A("One sheet per page", "ورقة لكل صفحة")),
              ("single", A("All tables in one sheet", "كل الجداول في ورقة واحدة"))]),
            ("borderless", A("Detect tables without borders", "اكتشاف الجداول بدون حدود"), "bool", True),
            ("text", A("Also export text outside tables", "تصدير النص خارج الجداول أيضاً"), "bool", False),
            ("numbers", A("Convert numbers, % and dates", "تحويل الأرقام والنسب والتواريخ"), "bool", True),
            ("ocr", A("OCR scanned pages", "OCR للصفحات الممسوحة"), "bool", True),
            ("pages", tr("Pages to process"), "text", "all"),
            ("out", A("Save to", "حفظ إلى"), "file_save", os.path.splitext(t.pdf.path or t.pdf.title)[0] + ".xlsx", "Excel (*.xlsx)")],
            width=540)
        if not v or not v["out"]:
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            st = pdfconv.pdf_to_xlsx(t.doc, v["out"], v["layout"], v["text"], True, v["borderless"], v["numbers"],
                                     ops.parse_ranges(v["pages"], t.doc.page_count), v["ocr"])
        finally:
            QApplication.restoreOverrideCursor()
        self.statusBar().showMessage(A(f"{st['tables']} tables, {st['rows']} rows exported",
                                       f"تم تصدير {st['tables']} جدول و {st['rows']} صف"), 6000)
        QDesktopServices.openUrl(QUrl.fromLocalFile(v["out"]))

    # ------------------------------------------------------------------ share
    def _email_file(self, path):
        sent = False
        if sys.platform.startswith("win"):
            try:
                import win32com.client
                mail = win32com.client.Dispatch("Outlook.Application").CreateItem(0)
                mail.Subject = os.path.basename(path)
                mail.Attachments.Add(os.path.abspath(path))
                mail.Display()
                sent = True
            except Exception:
                sent = False
        if not sent:
            QApplication.clipboard().setText(path)
            QDesktopServices.openUrl(QUrl(f"mailto:?subject={os.path.basename(path)}"))
            QDesktopServices.openUrl(QUrl.fromLocalFile(os.path.dirname(path)))
            self.statusBar().showMessage(tr("File path copied — attach it to the email"), 8000)

    # ------------------------------------------------------------------ accessibility
    def _access_issues(self, t):
        d = t.doc
        md = d.metadata or {}
        issues, ok = [], []
        (ok if md.get("title") else issues).append(A("Document title", "عنوان المستند"))
        cat = d.pdf_catalog()
        lang = d.xref_get_key(cat, "Lang")[1] if cat else ""
        (ok if lang not in ("", "null") else issues).append(A("Document language", "لغة المستند"))
        tagged = d.xref_get_key(cat, "MarkInfo")[0] != "null" if cat else False
        (ok if tagged else issues).append(A("Tagged PDF (structure)", "مستند معلَّم (بنية)"))
        (ok if d.get_toc() else issues).append(A("Bookmarks for navigation", "إشارات مرجعية للتنقل"))
        scanned = [i + 1 for i in range(d.page_count) if not d[i].get_text().strip() and d[i].get_images()]
        (issues if scanned else ok).append(A("Text is real text (not scanned images)", "النص حقيقي وليس صوراً ممسوحة") +
                                           (f": {scanned[:10]}" if scanned else ""))
        enc = md.get("encryption")
        (issues if enc and not d.permissions & fitz.PDF_PERM_ACCESSIBILITY else ok).append(
            A("Screen readers allowed", "السماح لقارئات الشاشة"))
        return ok, issues

    def do_access_report(self):
        t = self._need()
        if not t:
            return
        ok, issues = self._access_issues(t)
        txt = (A("Passed", "ناجح") + ":\n" + "\n".join("  ✔ " + s for s in ok) + "\n\n" +
               A("Needs attention", "يحتاج معالجة") + ":\n" + "\n".join("  ✖ " + s for s in issues))
        txt += "\n\n" + A("Tips: set title & language in Accessibility › Title & language; run OCR for scanned pages; "
                          "create bookmarks from headings.",
                          "نصائح: اضبط العنوان واللغة من إمكانية الوصول › العنوان واللغة، واستخدم OCR للصفحات الممسوحة، "
                          "وأنشئ إشارات من العناوين.")
        ReportDialog(self, A("Accessibility report", "تقرير إمكانية الوصول"), txt).exec()

    def do_access_check(self):
        t = self._need()
        if not t:
            return
        ok, issues = self._access_issues(t)
        QMessageBox.information(self, A("Accessibility check", "التحقق من إمكانية الوصول"),
                                A(f"{len(ok)} passed, {len(issues)} need attention.", f"{len(ok)} ناجح، {len(issues)} يحتاج معالجة."))

    # ------------------------------------------------------------------ bookmarks
    def _toc(self):
        t = self._need()
        return t, (t.doc.get_toc() if t else [])

    def _set_toc(self, t, toc, msg=""):
        t.pdf.checkpoint()
        t.doc.set_toc(toc)
        t.done(-1, msg)
        if not (t.stack.isVisible() and t.stack.currentIndex() == 1):
            t.show_panel(1)

    def do_bm_from_text(self):
        self.do_bookmarks_from_text()

    def do_bm_each(self):
        t, toc = self._toc()
        if not t:
            return
        v = ask(self, A("Bookmark every nth page", "إشارة لكل صفحة"), [("n", "n", "int", 1, (1, 1000)),
                                                                      ("fmt", A("Title", "العنوان"), "text", A("Page {page}", "صفحة {page}"))])
        if v:
            new = [[1, v["fmt"].replace("{page}", str(i + 1)), i + 1] for i in range(0, t.doc.page_count, v["n"])]
            self._set_toc(t, new)

    def do_bm_from_links(self):
        t, toc = self._toc()
        if not t:
            return
        new = list(toc)
        for i, pg in enumerate(t.doc):
            for l in pg.get_links():
                if l.get("kind") == fitz.LINK_GOTO and l.get("page", -1) >= 0:
                    title = pg.get_textbox(fitz.Rect(l["from"])).strip() or A(f"Page {l['page'] + 1}", f"صفحة {l['page'] + 1}")
                    new.append([1, title[:120], l["page"] + 1])
        self._set_toc(t, new)

    def do_bm_import_txt(self):
        t, toc = self._toc()
        if not t:
            return
        fn, _ = QFileDialog.getOpenFileName(self, A("From text file", "من ملف نصي"), "", "Text (*.txt)")
        if not fn:
            return
        new = []
        for line in open(fn, encoding="utf-8-sig"):
            if not line.strip():
                continue
            lvl = (len(line) - len(line.lstrip("\t"))) + 1
            m = re.match(r"(.*?)[\s.\-]+(\d+)\s*$", line.strip())
            if m:
                new.append([min(lvl, (new[-1][0] + 1) if new else 1), m.group(1), int(m.group(2))])
        self._set_toc(t, new)

    def do_bm_add_text(self):
        t, toc = self._toc()
        if not t or not toc:
            return
        v = ask(self, A("Add text", "أضف نصاً"), [("pre", A("Prefix", "بادئة"), "text", ""), ("suf", A("Suffix", "لاحقة"), "text", "")])
        if v:
            self._set_toc(t, [[l, v["pre"] + ti + v["suf"], p] for l, ti, p in toc])

    def do_bm_case(self):
        t, toc = self._toc()
        if not t or not toc:
            return
        v = ask(self, A("Change case", "تغيير حالة الأحرف"), [("c", "", "choice", "title", [("upper", "UPPER"), ("lower", "lower"),
                                                                                           ("title", "Title Case")])])
        if v:
            f = {"upper": str.upper, "lower": str.lower, "title": str.title}[v["c"]]
            self._set_toc(t, [[l, f(ti), p] for l, ti, p in toc])

    def do_bm_find_replace(self):
        t, toc = self._toc()
        if not t or not toc:
            return
        v = ask(self, A("Find & replace", "البحث والاستبدال"), [("f", A("Find", "بحث عن"), "text", ""),
                                                               ("r", A("Replace with", "استبدال بـ"), "text", "")])
        if v and v["f"]:
            self._set_toc(t, [[l, ti.replace(v["f"], v["r"]), p] for l, ti, p in toc])

    def do_bm_sort(self):
        t, toc = self._toc()
        if t and toc:
            self._set_toc(t, [[1, ti, p] for _, ti, p in sorted(toc, key=lambda e: (e[2], e[1]))])

    def do_bm_check(self):
        t, toc = self._toc()
        if not t:
            return
        bad = [e for e in toc if not (1 <= e[2] <= t.doc.page_count)]
        if bad and QMessageBox.question(self, A("Check", "التحقق"),
                                        A(f"{len(bad)} broken bookmarks. Remove them?", f"{len(bad)} إشارة معطلة. حذفها؟")) == QMessageBox.Yes:
            self._set_toc(t, [e for e in toc if e not in bad])
        elif not bad:
            QMessageBox.information(self, A("Check", "التحقق"), A("All bookmarks are valid.", "كل الإشارات سليمة."))

    def do_bm_merge_dups(self):
        t, toc = self._toc()
        if not t or not toc:
            return
        seen, new = set(), []
        for e in toc:
            k = (e[1].strip(), e[2])
            if k not in seen:
                seen.add(k)
                new.append(e)
        self._set_toc(t, new, A(f"{len(toc) - len(new)} duplicates removed", f"حُذف {len(toc) - len(new)} مكرر"))

    def do_sort_pages_bm(self):
        t, toc = self._toc()
        if not t or not toc:
            return
        tops = sorted([e for e in toc if e[0] == 1], key=lambda e: e[1].lower())
        starts = sorted({e[2] for e in toc if e[0] == 1})
        ranges = {}
        for k, s in enumerate(starts):
            ranges[s] = list(range(s - 1, (starts[k + 1] - 1) if k + 1 < len(starts) else t.doc.page_count))
        order = []
        for e in tops:
            order += [p for p in ranges.get(e[2], []) if p not in order]
        order += [p for p in range(t.doc.page_count) if p not in order]
        t.pdf.checkpoint()
        t.doc.select(order)
        t.done(-1)

    def do_toc_page(self):
        t, toc = self._toc()
        if not t or not toc:
            QMessageBox.information(self, A("Table of contents", "جدول المحتويات"), A("Create bookmarks first.", "أنشئ إشارات مرجعية أولاً."))
            return
        rtl = i18n.LANG == "ar" or any(ops.is_rtl(e[1]) for e in toc)
        title = A("Contents", "المحتويات")
        rows = "".join(f"<tr><td style='padding-{'right' if rtl else 'left'}:{(l - 1) * 18}px'>{_html.escape(ti)}</td>"
                       f"<td style='text-align:{'left' if rtl else 'right'}'>{p + 1}</td></tr>" for l, ti, p in toc)
        html = f"<div dir='{'rtl' if rtl else 'ltr'}'><h1>{title}</h1><table style='width:100%'>{rows}</table></div>"
        css = "* {font-family: sans-serif} h1 {font-size: 20pt} td {font-size: 11pt; padding: 3px 0} table {border-collapse: collapse}"
        w, h = t.doc[0].rect.width, t.doc[0].rect.height
        import io
        story = fitz.Story(html=html, user_css=css)
        buf = io.BytesIO()
        writer = fitz.DocumentWriter(buf)
        box = fitz.Rect(0, 0, w, h)
        more = True
        while more:
            dev = writer.begin_page(box)
            more, _ = story.place(box + (54, 54, -54, -54))
            story.draw(dev)
            writer.end_page()
        writer.close()
        tmp = fitz.open("pdf", buf.getvalue())
        n = tmp.page_count
        t.pdf.checkpoint()
        t.doc.insert_pdf(tmp, start_at=0)
        t.doc.set_toc([[1, title, 1]] + [[l, ti, p + n] for l, ti, p in toc])
        # clickable entries
        for i in range(n):
            pg = t.doc[i]
            for l, ti, p in toc:
                for r in pg.search_for(ti[:40])[:1]:
                    pg.insert_link({"kind": fitz.LINK_GOTO, "from": fitz.Rect(54, r.y0, w - 54, r.y1), "page": p + n - 1,
                                    "to": fitz.Point(0, 0)})
        t.done(-1)
        t.view.go_to(0)

    def _toc_export(self, html):
        t, toc = self._toc()
        if not t or not toc:
            return
        ext = "html" if html else "txt"
        fn, _ = QFileDialog.getSaveFileName(self, A("Export", "تصدير"), os.path.splitext(t.pdf.title)[0] + "_bookmarks." + ext, f"*.{ext}")
        if not fn:
            return
        if html:
            body = "".join(f"<div style='margin-inline-start:{(l - 1) * 20}px'>{_html.escape(ti)} — {p}</div>" for l, ti, p in toc)
            open(fn, "w", encoding="utf-8").write(f"<html><meta charset='utf-8'><body dir='auto'>{body}</body></html>")
        else:
            open(fn, "w", encoding="utf-8").write("\n".join("\t" * (l - 1) + f"{ti} {p}" for l, ti, p in toc))
        QDesktopServices.openUrl(QUrl.fromLocalFile(fn))

    def do_bm_export_html(self):
        self._toc_export(True)

    def do_bm_export_txt(self):
        self._toc_export(False)

    # ------------------------------------------------------------------ help
    def do_set_default(self):
        if sys.platform.startswith("win"):
            QDesktopServices.openUrl(QUrl("ms-settings:defaultapps"))
            QMessageBox.information(self, A("Set as default", "تعيين كبرنامج افتراضي"),
                                    A("Choose Waraq PDF for .pdf files in the Windows settings window.",
                                      "اختر Waraq PDF لملفات ‎.pdf‎ في نافذة إعدادات ويندوز."))
        else:
            QMessageBox.information(self, A("Set as default", "تعيين كبرنامج افتراضي"), A("Use your system settings.", "استخدم إعدادات النظام."))

    def do_cmdline_help(self):
        txt = ("WaraqPDF.exe [file1.pdf file2.pdf ...]\n    " + A("open files in the editor", "فتح الملفات في المحرر") +
               "\n\nWaraqPDF.exe --tools [files...]\n    " + A("open Batch tools with these files", "فتح الأدوات الدفعية مع هذه الملفات"))
        ReportDialog(self, A("Command line help", "مساعدة سطر الأوامر"), txt).exec()

    def do_report_problem(self):
        QDesktopServices.openUrl(QUrl("mailto:?subject=Waraq%20PDF%20-%20problem%20report"))
